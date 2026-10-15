import io
import unittest
import zipfile

try:
    import docx
    from pypdf import PdfReader, PdfWriter
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    from battle_test.web import extract
    from battle_test.web.extract import ExtractError, extract_text
except ImportError:  # web dependencies are optional
    extract = None


def make_docx(build) -> bytes:
    doc = docx.Document()
    build(doc)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def make_pdf(pages: list[list[str]]) -> bytes:
    """A PDF with each page's lines drawn down the page."""
    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=letter)
    for lines in pages:
        y = 740
        for line in lines:
            c.drawString(72, y, line)
            y -= 16
        c.showPage()
    c.save()
    return out.getvalue()


@unittest.skipIf(extract is None, "web dependencies not installed")
class ExtractTest(unittest.TestCase):
    def test_text_files_in_utf8_or_windows_encoding(self):
        self.assertEqual(extract_text("c.txt", "Café § 1\r\nline two".encode("utf-8-sig")), "Café § 1\nline two")
        self.assertEqual(extract_text("c.TXT", "Café".encode("cp1252")), "Café")

    def test_refuses_other_types_empty_and_oversized_files(self):
        for name, data, message in [
            ("c.doc", b"x", "save it as .docx"),
            ("complaint", b"x", "Upload a Word"),
            ("c.txt", b"", "empty"),
            ("c.txt", b"x" * (extract.MAX_UPLOAD_BYTES + 1), "too large"),
            ("c.txt", b"  \n\n ", "No text was found"),
            ("c.docx", b"not a zip", "isn't a readable Word"),
            ("c.pdf", b"not a pdf", "isn't a readable PDF"),
            ("c.pdf", b"%PDF-1.4 garbage", "couldn't be read"),
        ]:
            with self.subTest(name=name, message=message), self.assertRaisesRegex(ExtractError, message):
                extract_text(name, data)

    def test_word_keeps_order_tables_and_automatic_numbering(self):
        def build(doc):
            doc.add_paragraph("IN THE THIRD JUDICIAL DISTRICT COURT")
            table = doc.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "DANA WHITFIELD,\nPlaintiff,"
            table.rows[0].cells[1].text = "COMPLAINT"
            doc.add_paragraph("")
            for text in ("Plaintiff lives in Provo.", "Defendant is a roofer."):
                doc.add_paragraph(text, style="List Number")
            doc.add_paragraph("A bullet point.", style="List Bullet")
        text = extract_text("complaint.docx", make_docx(build))
        self.assertEqual(text.split("\n\n"), [
            "IN THE THIRD JUDICIAL DISTRICT COURT",
            "DANA WHITFIELD, Plaintiff, | COMPLAINT",
            "1. Plaintiff lives in Provo.",
            "2. Defendant is a roofer.",
            "• A bullet point.",
        ])

    def test_word_file_that_expands_too_far_is_refused(self):
        data = io.BytesIO()
        with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("word/document.xml", b"\0" * (extract.MAX_DOCX_EXPANDED_BYTES + 1))
        with self.assertRaisesRegex(ExtractError, "too large to read"):
            extract_text("bomb.docx", data.getvalue())

    def test_pdf_lines_are_joined_and_page_furniture_dropped(self):
        footer = "COMPLAINT - Whitfield v. Summit Peak Roofing"
        pages = [
            ["COMPLAINT", "1", "1. Plaintiff Dana Whitfield lives in Provo and con-", "tracted with Defendant",
             "for a new roof.", "2", "2. Defendant is a Utah company.", footer, "Page 1 of 3"],
            ["3. The claim is timely under Utah Code", "§ 78B-2-309.", footer, "Page 2 of 3"],
            ["PRAYER FOR RELIEF", "Plaintiff asks for damages.", footer, "Page 3 of 3"],
        ]
        text = extract_text("complaint.pdf", make_pdf(pages))
        self.assertEqual(text.split("\n\n"), [
            "COMPLAINT",
            "1. Plaintiff Dana Whitfield lives in Provo and contracted with Defendant for a new roof.",
            "2. Defendant is a Utah company.",
            "3. The claim is timely under Utah Code § 78B-2-309.",
            "PRAYER FOR RELIEF",
            "Plaintiff asks for damages.",
        ])

    def test_scanned_encrypted_and_overlong_pdfs_are_refused(self):
        out = io.BytesIO()
        c = canvas.Canvas(out)
        c.rect(72, 72, 200, 200, fill=1)  # a page with a picture and no text
        c.showPage()
        c.save()
        with self.assertRaisesRegex(ExtractError, "looks scanned"):
            extract_text("scan.pdf", out.getvalue())

        writer = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf([["Secret text."]]))))
        writer.encrypt("pw")
        locked = io.BytesIO()
        writer.write(locked)
        with self.assertRaisesRegex(ExtractError, "password-protected"):
            extract_text("locked.pdf", locked.getvalue())

        long_pdf = make_pdf([["x"]] * (extract.MAX_PDF_PAGES + 1))
        with self.assertRaisesRegex(ExtractError, "pages: the limit is"):
            extract_text("long.pdf", long_pdf)


if __name__ == "__main__":
    unittest.main()
