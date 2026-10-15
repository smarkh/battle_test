import io
import re
import tempfile
import unittest
from pathlib import Path

try:
    import pyarrow as pa
    import pyarrow.parquet as pq
    from fastapi.testclient import TestClient

    from battle_test.config import CorpusConfig, Plan, Plans
    from battle_test.corpus import build_index
    from battle_test.web.app import case_title, create_app, facts_text, safe_next
    from battle_test.web.demo import DemoClient
    from battle_test.web import export, render
    from battle_test.web.render import document_html
except ImportError:  # web and corpus-build dependencies are optional
    TestClient = None


def _row(act_id, citation, title, text):
    return {"act_id": act_id, "citation": citation, "document_type": "statute", "title_name": "Code",
            "chapter_name": "None", "section_title": title, "display_path": citation,
            "act_status": "in_force", "text": text, "source_url": f"https://le.utah.gov/{act_id}",
            "last_amended_year": 2020}


FACTS = {"state": "UT", "mode": "facts", "rounds": "1", "plaintiff": "Dana Whitfield",
         "defendant": "Summit Peak Roofing LLC", "timeline": "Paid a deposit; roofer never came back."}

# Throwaway test credentials only.
PW = "correct horse battery"
NEW_PW = "a brand new passphrase"


@unittest.skipIf(TestClient is None, "web or pyarrow dependencies not installed")
class WebAppTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        pq.write_table(pa.Table.from_pylist([
            _row("R56", "Utah R. Civ. P. 56", "Rule 56. Summary judgment", "A party may move for summary judgment."),
            _row("L1", "Utah Code § 78B-2-309", "Within six years", "An action upon a contract in writing."),
        ]), root / "us_ut_court_rules.parquet")
        cfg = CorpusConfig("vTEST", "http://unused", root, ("ut",), ("court_rules",))
        build_index(cfg, {"snapshot_date": "2026-08-14"}, [root / "us_ut_court_rules.parquet"])
        plans = Plans({"unlimited": Plan("unlimited", "Unlimited", None, "month"),
                       "two": Plan("two", "Two a month", 2, "month"),
                       "trial": Plan("trial", "Trial", 1, "total")}, default="unlimited")
        self.app = create_app(client=DemoClient(delay_per_word=0), model_label="demo",
                              data_dir=root / "cases", law_db=cfg.db_path, plans=plans)
        self.app.state.auth.create_user("dana", PW)
        self.app.state.auth.create_user("other", PW)
        self.client = TestClient(self.app)
        self.client.__enter__()  # runs the lifespan, which starts the worker
        self.login("dana")

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.tmp.cleanup()

    def login(self, username, password=PW, client=None):
        client = client or self.client
        return client.post("/login", data={"username": username, "password": password, "next": "/"},
                           follow_redirects=False)

    def csrf(self, client=None):
        page = (client or self.client).get("/cases/new").text
        return re.search(r'name="csrf" value="([^"]+)"', page).group(1)

    def submit(self, client=None, **overrides):
        client = client or self.client
        data = {**FACTS, "csrf": self.csrf(client), **overrides}
        return client.post("/cases", data=data, follow_redirects=False)

    def finished_case(self, **overrides):
        case_url = self.submit(**overrides).headers["location"]
        self.assertTrue(self.app.state.worker.wait_idle())
        return case_url

    # --- signing in ---------------------------------------------------------

    def test_pages_need_sign_in(self):
        anon = TestClient(self.app)
        for path in ("/", "/cases/new", "/account", "/cases/abc", "/cases/abc/events", "/cases/abc/download.md"):
            with self.subTest(path=path):
                response = anon.get(path, follow_redirects=False)
                self.assertEqual(response.status_code, 303)
                self.assertTrue(response.headers["location"].startswith("/login?next="))
        self.assertEqual(anon.post("/cases", data=FACTS, follow_redirects=False).status_code, 303)
        self.assertEqual(anon.get("/login").status_code, 200)

    def test_wrong_password_is_refused(self):
        anon = TestClient(self.app)
        response = self.login("dana", "not the password", client=anon)
        self.assertEqual(response.status_code, 401)
        self.assertIn("Wrong username or password.", response.text)
        self.assertEqual(anon.get("/", follow_redirects=False).status_code, 303)

    def test_session_cookie_is_httponly_and_samesite(self):
        cookie = self.login("dana", client=TestClient(self.app)).headers["set-cookie"]
        self.assertIn("bt_session=", cookie)
        self.assertIn("HttpOnly", cookie)
        self.assertIn("SameSite=lax", cookie)

    def test_login_redirects_only_within_the_site(self):
        for target, expected in [("/cases/new", "/cases/new"), ("//evil.example", "/"),
                                 ("https://evil.example", "/"), ("/\\evil.example", "/")]:
            with self.subTest(target=target):
                self.assertEqual(safe_next(target), expected)
        response = TestClient(self.app).post(
            "/login", data={"username": "dana", "password": PW, "next": "//evil.example"}, follow_redirects=False)
        self.assertEqual(response.headers["location"], "/")

    def test_logout(self):
        token = re.search(r'name="csrf" value="([^"]+)"', self.client.get("/").text).group(1)
        self.client.post("/logout", data={"csrf": token}, follow_redirects=False)
        self.assertEqual(self.client.get("/", follow_redirects=False).status_code, 303)

    def test_change_password(self):
        response = self.client.post("/account", data={"csrf": self.csrf(), "current": PW, "new": NEW_PW,
                                                     "repeat": NEW_PW})
        self.assertIn("Password changed", response.text)
        self.assertEqual(self.client.get("/").status_code, 200)  # this browser stays signed in
        self.assertEqual(self.login("dana", PW, client=TestClient(self.app)).status_code, 401)
        self.assertEqual(self.login("dana", NEW_PW, client=TestClient(self.app)).status_code, 303)

    def test_change_password_needs_the_current_one(self):
        response = self.client.post("/account", data={"csrf": self.csrf(), "current": "wrong wrong wrong",
                                                     "new": NEW_PW, "repeat": NEW_PW})
        self.assertEqual(response.status_code, 400)
        self.assertIn("current password is wrong", response.text)

    def test_forms_need_csrf_token_and_same_origin(self):
        self.assertEqual(self.submit(csrf="wrong").status_code, 403)
        cross = self.client.post("/cases", data={**FACTS, "csrf": self.csrf()},
                                 headers={"origin": "https://evil.example"}, follow_redirects=False)
        self.assertEqual(cross.status_code, 403)
        self.assertIn("Start your first case", self.client.get("/").text)

    def test_browser_style_origin_headers(self):
        # Browsers send Origin on form posts. Same-site is fine; "null" (what
        # a no-referrer policy produces) and other sites are refused.
        anon = TestClient(self.app)
        form = {"username": "dana", "password": PW, "next": "/"}
        ok = anon.post("/login", data=form, headers={"origin": "http://testserver"}, follow_redirects=False)
        self.assertEqual(ok.status_code, 303)
        for origin in ("null", "http://evil.example"):
            with self.subTest(origin=origin):
                refused = TestClient(self.app).post("/login", data=form, headers={"origin": origin},
                                                    follow_redirects=False)
                self.assertEqual(refused.status_code, 403)
        self.assertEqual(self.client.get("/").headers["referrer-policy"], "same-origin")

    def test_security_headers(self):
        headers = self.client.get("/").headers
        self.assertIn("default-src 'self'", headers["content-security-policy"])
        self.assertEqual(headers["x-frame-options"], "DENY")
        self.assertEqual(headers["cache-control"], "no-store")

    def test_static_files_are_linked_with_a_version(self):
        page = self.client.get("/cases/new").text
        css, js = re.findall(r'"(/static/app\.(?:css|js)\?v=[0-9a-f]{12})"', page)
        self.assertEqual(css.split("?v=")[1], js.split("?v=")[1])
        self.assertEqual(self.client.get(css).status_code, 200)
        self.assertIn("EventSource", self.client.get(js).text)

    # --- setup links ----------------------------------------------------------

    def test_setup_link_sets_the_password_once(self):
        auth = self.app.state.auth
        auth.create_user("newbie", None, plan="solo")
        token = auth.create_setup_token("newbie")
        anon = TestClient(self.app)
        self.assertEqual(self.login("newbie", PW, client=anon).status_code, 401)  # no password yet

        page = anon.get(f"/setup/{token}")
        self.assertEqual(page.status_code, 200)
        self.assertIn("Your username is <strong>newbie</strong>", page.text)
        self.assertEqual(page.headers["cache-control"], "no-store")
        self.assertNotIn("Sign out", page.text)

        mismatch = anon.post(f"/setup/{token}", data={"new": NEW_PW, "repeat": PW})
        self.assertEqual(mismatch.status_code, 400)
        self.assertIn("didn't match", mismatch.text.replace("&#39;", "'"))
        short = anon.post(f"/setup/{token}", data={"new": "short", "repeat": "short"})
        self.assertEqual(short.status_code, 400)
        self.assertIn("at least 12", short.text)
        cross = anon.post(f"/setup/{token}", data={"new": NEW_PW, "repeat": NEW_PW},
                          headers={"origin": "https://evil.example"})
        self.assertEqual(cross.status_code, 403)

        done = anon.post(f"/setup/{token}", data={"new": NEW_PW, "repeat": NEW_PW}, follow_redirects=False)
        self.assertEqual(done.headers["location"], "/login?set=1")
        self.assertIn("Password set.", anon.get("/login?set=1").text)
        self.assertEqual(self.login("newbie", NEW_PW, client=anon).status_code, 303)

        # The link is now dead, for viewing and for posting.
        again = TestClient(self.app)
        self.assertEqual(again.get(f"/setup/{token}").status_code, 404)
        reused = again.post(f"/setup/{token}", data={"new": PW, "repeat": PW})
        self.assertEqual(reused.status_code, 404)
        self.assertIn("expired or was already used", reused.text)
        self.assertEqual(self.login("newbie", PW, client=again).status_code, 401)

    def test_unknown_setup_link_reveals_nothing(self):
        response = TestClient(self.app).get("/setup/not-a-real-token")
        self.assertEqual(response.status_code, 404)
        self.assertIn("expired or was already used", response.text)
        self.assertNotIn("<form", response.text)

    def test_setup_link_signs_this_browser_out_of_another_account(self):
        token = self.app.state.auth.create_setup_token("other")
        self.client.post(f"/setup/{token}", data={"new": NEW_PW, "repeat": NEW_PW}, follow_redirects=False)
        self.assertEqual(self.client.get("/", follow_redirects=False).status_code, 303)  # dana was signed out here

    # --- admin pages -----------------------------------------------------------

    def admin_client(self):
        self.app.state.auth.create_user("boss", PW, is_admin=True)
        boss = TestClient(self.app)
        self.login("boss", client=boss)
        return boss

    def admin_post(self, boss, path, password=PW, **data):
        return boss.post(path, data={"csrf": self.csrf(boss), "admin_password": password, **data})

    def events(self, subject=None):
        return [(e.kind, e.subject, e.actor, e.detail) for e in self.app.state.auth.events(subject)]

    def test_admin_pages_are_for_admins_only(self):
        boss = self.admin_client()
        paths = ("/admin", "/admin/activity", "/admin/users/dana")
        for path in paths:
            with self.subTest(path=path):
                self.assertEqual(boss.get(path).status_code, 200)
                self.assertEqual(self.client.get(path).status_code, 404)  # dana isn't an admin
                anon = TestClient(self.app).get(path, follow_redirects=False)
                self.assertEqual(anon.status_code, 303)
        for path in ("/admin/users", "/admin/users/other/edit", "/admin/users/other/invite",
                     "/admin/users/other/disable", "/admin/users/other/enable"):
            with self.subTest(path=path):
                response = self.client.post(path, data={"csrf": self.csrf(), "admin_password": PW,
                                                        "username": "sneaky", "plan": "unlimited"})
                self.assertEqual(response.status_code, 404)
        self.assertIsNone(self.app.state.auth.get_user("sneaky"))
        self.assertFalse(self.app.state.auth.get_user("other").disabled)
        self.assertIn('<a href="/admin">Admin</a>', boss.get("/").text)
        self.assertNotIn("/admin", self.client.get("/").text)
        self.assertEqual(boss.get("/admin/users/nobody").status_code, 404)

    def test_admin_adds_an_account_and_gets_its_setup_link(self):
        boss = self.admin_client()
        response = self.admin_post(boss, "/admin/users", username="Newbie", plan="two", firm="Acme",
                                   paid_through="2099-12-31")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Created newbie.", response.text)
        token = re.search(r'<code class="setup-link">http://127\.0\.0\.1:8000/setup/([^<]+)</code>',
                          response.text).group(1)
        user = self.app.state.auth.setup_user(token)
        self.assertEqual((user.username, user.plan, user.firm, user.paid_through, user.needs_setup, user.is_admin),
                         ("newbie", "two", "acme", "2099-12-31", True, False))
        # The link is shown once: it isn't on the page when it's opened again.
        self.assertNotIn("setup-link", boss.get("/admin/users/newbie").text)
        self.assertIn("no password yet", boss.get("/admin").text)
        self.assertEqual([e[:3] for e in self.events("newbie")],
                         [("setup link issued", "newbie", "boss"), ("paid-through changed", "newbie", "boss"),
                          ("account created", "newbie", "boss")])

        duplicate = self.admin_post(boss, "/admin/users", username="newbie", plan="two")
        self.assertEqual(duplicate.status_code, 400)
        self.assertIn("already exists", duplicate.text)
        bad_plan = self.admin_post(boss, "/admin/users", username="erin", plan="platinum")
        self.assertEqual(bad_plan.status_code, 400)
        self.assertIsNone(self.app.state.auth.get_user("erin"))

    def test_admin_changes_need_the_admins_password_and_csrf(self):
        boss = self.admin_client()
        for path, data in (("/admin/users", {"username": "erin", "plan": "two"}),
                           ("/admin/users/other/edit", {"plan": "two", "firm": "", "paid_through": ""}),
                           ("/admin/users/other/invite", {}), ("/admin/users/other/disable", {})):
            with self.subTest(path=path):
                refused = self.admin_post(boss, path, password="not the password", **data)
                self.assertEqual(refused.status_code, 403)
                self.assertIn("Your password is wrong, so nothing was changed.", refused.text)
                self.assertNotIn("setup-link", refused.text)
                no_csrf = boss.post(path, data={"csrf": "wrong", "admin_password": PW, **data})
                self.assertEqual(no_csrf.status_code, 403)
        other = self.app.state.auth.get_user("other")
        self.assertEqual((other.plan, other.disabled), ("", False))
        self.assertIsNone(self.app.state.auth.get_user("erin"))
        self.assertEqual([e[0] for e in self.events("other")], ["account created"])  # from setUp only
        self.assertEqual({e[0] for e in self.events("boss")},
                         {"signed in", "admin password refused", "account created"})

    def test_admin_edits_plan_firm_and_paid_through(self):
        boss = self.admin_client()
        response = self.admin_post(boss, "/admin/users/dana/edit", plan="two", firm="Acme", paid_through="2026-01-31")
        self.assertIn("Saved.", response.text)
        dana = self.app.state.auth.get_user("dana")
        self.assertEqual((dana.plan, dana.firm, dana.paid_through), ("two", "acme", "2026-01-31"))
        self.assertIn("expired 2026-01-31", response.text)
        self.assertIn("ended on 2026-01-31", self.client.get("/cases/new").text)  # it applies at once

        # Saving again with one change logs just that change.
        before = len(self.events("dana"))
        self.admin_post(boss, "/admin/users/dana/edit", plan="two", firm="acme", paid_through="")
        self.assertEqual(self.events("dana")[0], ("paid-through changed", "dana", "boss", "2026-01-31 to no end"))
        self.assertEqual(len(self.events("dana")), before + 1)

        bad = self.admin_post(boss, "/admin/users/dana/edit", plan="two", firm="two words", paid_through="")
        self.assertEqual(bad.status_code, 400)
        self.assertIn("Firm names", bad.text)

    def test_admin_disables_enables_and_issues_setup_links(self):
        boss = self.admin_client()
        response = self.admin_post(boss, "/admin/users/dana/invite")
        token = re.search(r"/setup/([^<]+)</code>", response.text).group(1)
        self.assertEqual(self.app.state.auth.setup_user(token).username, "dana")

        response = self.admin_post(boss, "/admin/users/dana/disable")
        self.assertIn("Disabled, and signed out everywhere.", response.text)
        self.assertEqual(self.client.get("/", follow_redirects=False).status_code, 303)  # dana is signed out
        self.assertIsNone(self.app.state.auth.setup_user(token))
        self.assertEqual(self.admin_post(boss, "/admin/users/dana/invite").status_code, 400)
        self.admin_post(boss, "/admin/users/dana/enable")
        self.assertEqual(self.login("dana").status_code, 303)
        self.assertEqual([e[0] for e in self.events("dana")][:4],
                         ["signed in", "enabled", "disabled", "setup link issued"])

    def test_admin_accounts_are_managed_from_the_command_line(self):
        boss = self.admin_client()
        self.app.state.auth.create_user("chief", PW, is_admin=True)
        for action in ("invite", "disable", "enable"):
            with self.subTest(action=action):
                refused = self.admin_post(boss, f"/admin/users/chief/{action}")
                self.assertEqual(refused.status_code, 403)
                self.assertIn("done from the command line for admin accounts", refused.text)
                self.assertNotIn("setup-link", refused.text)
        self.assertFalse(self.app.state.auth.get_user("chief").disabled)
        # Nothing in the portal makes an admin.
        self.admin_post(boss, "/admin/users", username="erin", plan="two", is_admin="1", admin="1")
        self.assertFalse(self.app.state.auth.get_user("erin").is_admin)

    def test_admin_sees_usage_but_no_case_titles_or_text(self):
        case_url = self.finished_case()
        boss = self.admin_client()
        page = boss.get("/admin/users/dana").text
        self.assertIn("Complaint and motion drafted", page)
        self.assertRegex(page, r"1 of no limit cases used")
        for private in ("Whitfield", "Summit Peak", "roofer"):
            self.assertNotIn(private, page)
            self.assertNotIn(private, boss.get("/admin").text)
            self.assertNotIn(private, boss.get("/admin/activity").text)
        for path in (case_url, f"{case_url}/download.md", f"{case_url}/events"):
            self.assertEqual(boss.get(path).status_code, 404)  # an admin is no exception to case privacy

    def test_activity_log_records_sign_ins_without_mistyped_names(self):
        anon = TestClient(self.app)
        self.login("dana", "not the password", client=anon)
        self.login("my-actual-passphrase-typed-in-the-wrong-box", "x", client=anon)
        self.client.post("/account", data={"csrf": self.csrf(), "current": PW, "new": NEW_PW, "repeat": NEW_PW})
        self.assertEqual([e[:3] for e in self.events()][:3],
                         [("password changed", "dana", "dana"), ("sign-in failed", "dana", ""),
                          ("signed in", "dana", "dana")])
        self.assertIn("from testclient", self.events()[1][3])
        raw = (Path(self.tmp.name) / "cases" / "users.sqlite").read_bytes()
        self.assertNotIn(b"my-actual-passphrase", raw)
        boss = self.admin_client()
        log = boss.get("/admin/activity").text
        self.assertIn("sign-in failed", log)
        self.assertIn("password changed", log)
        self.assertRegex(boss.get("/admin").text, r"(?s)dana</a>.*?\d{4}-\d{2}-\d{2} \d{2}:\d{2}")

    # --- plans ------------------------------------------------------------------

    def use_plan(self, username, plan):
        self.app.state.auth.set_plan(username, plan)

    def test_case_allowance_is_shown_and_enforced(self):
        self.use_plan("dana", "two")
        self.assertRegex(self.client.get("/cases/new").text, r"2 of 2 cases left\s+this month\s+\(Two a month\)\.")
        self.finished_case()
        self.finished_case()
        account = self.client.get("/account").text
        self.assertIn("<strong>Two a month</strong>", account)
        self.assertRegex(account, r"2 of 2 cases used\s+this month")
        self.assertRegex(account, r"starts again on \d{4}-\d{2}-01")

        new = self.client.get("/cases/new").text
        self.assertIn("used all 2 cases in your Two a month plan this month", new)
        self.assertRegex(new, r'<button type="submit" class="button primary" disabled>')

        refused = self.submit(plaintiff="Kept Name")
        self.assertEqual(refused.status_code, 403)
        self.assertIn("used all 2 cases", refused.text)
        self.assertIn("Kept Name", refused.text)
        self.assertEqual(len(self.app.state.store.list()), 2)

        # Deleting a finished case doesn't give the use back.
        case_url = f"/cases/{self.app.state.store.list()[0].id}"
        self.client.post(f"{case_url}/delete", data={"csrf": self.csrf()})
        self.assertEqual(self.submit().status_code, 403)

        # Someone else is unaffected, and a bigger plan lifts the limit at once.
        other = TestClient(self.app)
        self.login("other", client=other)
        self.assertEqual(self.submit(client=other).status_code, 303)
        self.use_plan("dana", "unlimited")
        self.assertEqual(self.submit().status_code, 303)
        self.assertTrue(self.app.state.worker.wait_idle())
        self.assertNotIn("cases left", self.client.get("/cases/new").text)
        self.assertIn("no limit on cases", self.client.get("/account").text)

    def test_trial_allowance_never_renews(self):
        self.use_plan("dana", "trial")
        job = self.app.state.store.create(self.app.state.auth.get_user("dana").id, "UT", "facts", "x", 1, "t")
        self.app.state.store.update(job.id, status="done")
        with self.app.state.store._db:  # started long ago: still counts against a total
            self.app.state.store._db.execute("UPDATE usage SET created_at = '2025-01-01T00:00:00'")
        new = self.client.get("/cases/new").text
        self.assertIn("used all 1 cases in your Trial plan.", new)
        self.assertEqual(self.submit().status_code, 403)

    def test_failed_run_does_not_use_up_the_allowance(self):
        self.use_plan("dana", "trial")
        self.app.state.worker.runner = lambda *args: 1 / 0
        self.finished_case()
        self.assertEqual(self.app.state.store.list()[0].status, "failed")
        self.assertRegex(self.client.get("/cases/new").text, r"1 of 1 cases left\s+in your plan\s+\(Trial\)\.")

    def test_cap_on_cases_waiting_or_running(self):
        store, dana = self.app.state.store, self.app.state.auth.get_user("dana").id
        waiting = [store.create(dana, "UT", "facts", "x", 1, "t") for _ in range(3)]  # never run: not submitted
        refused = self.submit(plaintiff="Kept Name")
        self.assertEqual(refused.status_code, 429)
        self.assertIn("You already have 3 cases waiting or running", refused.text)
        self.assertIn("Kept Name", refused.text)
        self.assertEqual(len(store.list(dana)), 3)

        other = TestClient(self.app)
        self.login("other", client=other)
        self.assertEqual(self.submit(client=other).status_code, 303)  # someone else isn't held up
        store.update(waiting[0].id, status="done")
        self.assertEqual(self.submit().status_code, 303)
        for job in waiting[1:]:
            store.update(job.id, status="failed")  # let the worker skip them
        self.assertTrue(self.app.state.worker.wait_idle())

    def test_unpaid_plan_blocks_new_cases_but_not_existing_ones(self):
        auth = self.app.state.auth
        case_url = self.finished_case()
        auth.set_paid_through("dana", "2099-12-31")
        self.assertIn("Paid through 2099-12-31.", self.client.get("/account").text)
        self.assertEqual(self.submit().status_code, 303)
        self.assertTrue(self.app.state.worker.wait_idle())

        auth.set_paid_through("dana", "2026-01-31")
        self.assertEqual(self.client.get("/").status_code, 200)  # still signed in
        new = self.client.get("/cases/new").text
        self.assertIn("Your Unlimited plan ended on 2026-01-31", new)
        self.assertRegex(new, r'<button type="submit" class="button primary" disabled>')
        self.assertIn("ended on 2026-01-31", self.client.get("/account").text)
        refused = self.submit(plaintiff="Kept Name")
        self.assertEqual(refused.status_code, 403)
        self.assertIn("Kept Name", refused.text)
        self.assertEqual(len(self.app.state.store.list()), 2)

        # What they already have stays reachable, and deletable.
        self.assertIn("Citation check", self.client.get(case_url).text)
        self.assertEqual(self.client.get(f"{case_url}/download.md").status_code, 200)
        deleted = self.client.post(f"{case_url}/delete", data={"csrf": self.csrf()}, follow_redirects=False)
        self.assertEqual(deleted.status_code, 303)

        auth.set_paid_through("dana", "")  # renewed with no end date
        self.assertEqual(self.submit().status_code, 303)
        self.assertTrue(self.app.state.worker.wait_idle())

    def test_unpaid_firm_seat_stops_adding_to_the_shared_allowance(self):
        auth = self.app.state.auth
        for name in ("dana", "other"):
            self.use_plan(name, "two")
            auth.set_firm(name, "acme")
        auth.set_paid_through("other", "2026-01-31")
        self.assertRegex(self.client.get("/cases/new").text,
                         r"2 of 2 cases left\s+this month\s+\(shared by your firm's 1 seats\)")
        other = TestClient(self.app)
        self.login("other", client=other)
        self.assertEqual(self.submit(client=other).status_code, 403)
        self.assertEqual(self.submit().status_code, 303)  # the paid-up seat carries on
        self.assertTrue(self.app.state.worker.wait_idle())

    def test_firm_seats_share_one_allowance(self):
        auth = self.app.state.auth
        for name in ("dana", "other"):
            self.use_plan(name, "two")
            auth.set_firm(name, "Whitfield-Law")
        other = TestClient(self.app)
        self.login("other", client=other)
        self.assertRegex(self.client.get("/cases/new").text,
                         r"4 of 4 cases left\s+this month\s+\(shared by your firm's 2 seats\)\.")

        # One seat may use more than its own share...
        for _ in range(3):
            self.finished_case()
        self.assertRegex(other.get("/cases/new").text, r"1 of 4 cases left")
        account = other.get("/account").text
        self.assertRegex(account, r"3 of 4 cases used\s+this month")
        self.assertIn("shared by the 2 seats of your firm", account)
        self.assertIn("<strong>whitfield-law</strong>", account)
        # ...but the firm's total is the limit, for every seat.
        self.assertEqual(self.submit(client=other).status_code, 303)
        self.assertTrue(self.app.state.worker.wait_idle())
        for client in (self.client, other):
            refused = self.submit(client=client)
            self.assertEqual(refused.status_code, 403)
            self.assertIn("Your firm has used all 4 cases in its shared plan this month", refused.text)

        # Sharing an allowance doesn't share cases.
        self.assertEqual(len(self.app.state.store.list(auth.get_user("other").id)), 1)
        case_url = f"/cases/{self.app.state.store.list(auth.get_user('dana').id)[0].id}"
        self.assertEqual(other.get(case_url).status_code, 404)

        # A disabled account stops being a seat, but what it used still counts.
        auth.set_disabled("other", True)
        self.assertIn("Your firm has used all 2 cases", self.client.get("/cases/new").text)
        auth.set_disabled("other", False)

        # Leaving the firm: each account is back on its own plan and its own count.
        auth.set_firm("other", "")
        self.assertRegex(self.client.get("/account").text, r"3 of 2 cases used")
        self.login("other", client=other)
        self.assertRegex(other.get("/cases/new").text, r"1 of 2 cases left\s+this month\s+\(Two a month\)")

    def test_firm_with_an_unlimited_seat_has_no_limit(self):
        self.use_plan("dana", "trial")
        for name in ("dana", "other"):
            self.app.state.auth.set_firm(name, "acme")
        self.assertIn("no limit on cases", self.client.get("/account").text)
        self.finished_case()
        self.assertEqual(self.submit().status_code, 303)
        self.assertTrue(self.app.state.worker.wait_idle())

    def test_account_with_no_plan_gets_the_default(self):
        self.assertEqual(self.app.state.auth.get_user("dana").plan, "")
        self.assertIn("no limit on cases", self.client.get("/account").text)
        self.use_plan("dana", "a plan since removed from the config")
        self.assertIn("no limit on cases", self.client.get("/account").text)

    # --- cases ----------------------------------------------------------------

    def test_pages_render(self):
        for path in ("/", "/cases/new", "/account"):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200)
                self.assertIn("Sign out", response.text)
        self.assertIn("Start your first case", self.client.get("/").text)
        new = self.client.get("/cases/new").text
        self.assertIn("Draft the complaint and the motion", new)
        self.assertIn("Use my complaint, draft only the motion", new)

    def test_empty_case_list_shows_the_form(self):
        home = self.client.get("/").text
        self.assertIn("<h1>Start your first case</h1>", home)
        self.assertNotIn("Start one", home)
        self.assertIn('<form method="post" action="/cases" class="case-form" enctype="multipart/form-data">', home)
        self.assertIn("Draft the complaint and the motion", home)
        self.assertIn("Use my complaint, draft only the motion", home)
        self.assertNotIn('<table class="cases">', home)

        # The form on the home page works like the one on /cases/new.
        token = re.search(r'name="csrf" value="([^"]+)"', home).group(1)
        response = self.client.post("/cases", data={**FACTS, "csrf": token}, follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        self.assertTrue(self.app.state.worker.wait_idle())

        # Once there's a case, the home page is the case list again.
        home = self.client.get("/").text
        self.assertIn("<h1>Cases</h1>", home)
        self.assertIn('<table class="cases">', home)
        self.assertNotIn('action="/cases" class="case-form"', home)

    def test_full_run_from_facts(self):
        case_url = self.finished_case()
        page = self.client.get(case_url).text
        self.assertIn("Citation check", page)
        self.assertIn("current as of\n  <strong>2026-08-14</strong>", page)
        # The demo cites the first authority it's given (✅, linked) and an invented section (❌, flagged).
        self.assertRegex(page, r'<a class="cite ok" href="https://le\.utah\.gov/\w+"')
        self.assertIn('<mark class="flag">[⚠ NOT FOUND', page)
        self.assertIn('<mark class="placeholder">[CITATION NEEDED', page)
        self.assertIn("Opposition to Motion for Summary Judgment", page)
        self.assertNotIn("Reply in Support", page)  # 1 round
        self.assertIn("plaintiff model <code>demo</code>", page)
        self.assertNotIn("qwen", page)

        download = self.client.get(f"{case_url}/download.md")
        self.assertEqual(download.status_code, 200)
        self.assertIn("# Battle test — Utah", download.text)

        listing = self.client.get("/").text
        self.assertIn("Dana Whitfield v. Summit Peak Roofing LLC", listing)
        self.assertIn("Complaint and motion drafted", listing)
        self.assertIn('class="status done"', listing)

    def test_pasted_complaint_is_used_as_is(self):
        complaint = "COMPLAINT\nPlaintiff sues under Utah Code § 78B-2-309."
        case_url = self.finished_case(mode="complaint", complaint=complaint)
        page = self.client.get(case_url).text
        self.assertIn("user-supplied", page)
        self.assertIn("Plaintiff sues under", page)
        self.assertIn("Your complaint, motion drafted", self.client.get("/").text)

    def test_uploaded_complaint_is_shown_for_review_before_anything_runs(self):
        user_id = self.app.state.auth.get_user("dana").id
        text = "COMPLAINT\r\nPlaintiff sues under Utah Code § 78B-2-309 & more."
        data = {**FACTS, "mode": "complaint", "complaint": "old pasted text", "csrf": self.csrf()}
        response = self.client.post("/cases", data=data, follow_redirects=False,
                                    files={"complaint_file": ("my complaint.txt", text.encode(), "text/plain")})
        self.assertEqual(response.status_code, 200)
        self.assertIn("The text of my complaint.txt is below", response.text)
        self.assertIn("COMPLAINT\nPlaintiff sues under Utah Code § 78B-2-309 &amp; more.</textarea>", response.text)
        self.assertNotIn("old pasted text", response.text)  # the upload replaces what was pasted
        self.assertIn('enctype="multipart/form-data"', response.text)
        self.assertEqual(self.app.state.store.list(user_id), [])  # no case yet, no allowance used

        # Starting the run from the reviewed text works like a paste.
        case_url = self.submit(mode="complaint", complaint="COMPLAINT\nPlaintiff sues.").headers["location"]
        self.assertTrue(self.app.state.worker.wait_idle())
        self.assertIn("Plaintiff sues.", self.client.get(case_url).text)

    def test_unreadable_upload_explains_and_keeps_the_form(self):
        data = {**FACTS, "mode": "complaint", "complaint": "kept", "csrf": self.csrf()}
        response = self.client.post("/cases", data=data,
                                    files={"complaint_file": ("old.doc", b"\xd0\xcf\x11\xe0", "application/msword")})
        self.assertEqual(response.status_code, 200)
        self.assertIn("old.doc: Upload a Word (.docx), PDF (.pdf) or text (.txt) file.", response.text)
        self.assertIn(">kept</textarea>", response.text)

    def test_a_file_is_ignored_when_drafting_from_case_information(self):
        data = {**FACTS, "csrf": self.csrf()}
        response = self.client.post("/cases", data=data, follow_redirects=False,
                                    files={"complaint_file": ("c.txt", b"ignored", "text/plain")})
        self.assertEqual(response.status_code, 303)
        # And an empty file input, as a browser sends one, is no upload at all.
        data = {**FACTS, "mode": "complaint", "complaint": "COMPLAINT\nPasted.", "csrf": self.csrf()}
        response = self.client.post("/cases", data=data, follow_redirects=False,
                                    files={"complaint_file": ("", b"", "application/octet-stream")})
        self.assertEqual(response.status_code, 303)
        self.assertTrue(self.app.state.worker.wait_idle())  # both cases run before the test's teardown

    def test_cases_are_private_to_their_owner(self):
        case_url = self.finished_case()
        other = TestClient(self.app)
        self.login("other", client=other)
        for path in (case_url, f"{case_url}/events", f"{case_url}/download.md", f"{case_url}/download.docx",
                     f"{case_url}/download.pdf"):
            with self.subTest(path=path):
                self.assertEqual(other.get(path).status_code, 404)
        self.assertIn("Start your first case", other.get("/").text)

    def test_download_as_word_and_pdf(self):
        case_url = self.finished_case()
        page = self.client.get(case_url).text
        for fmt in ("docx", "pdf", "md"):
            self.assertIn(f'href="{case_url}/download.{fmt}"', page)

        word = self.client.get(f"{case_url}/download.docx")
        self.assertEqual(word.status_code, 200)
        self.assertEqual(word.headers["content-type"], export.FORMATS["docx"])
        self.assertRegex(word.headers["content-disposition"], r'attachment; filename="battle-test-ut-[\d-]+\.docx"')
        import docx
        text = "\n".join(p.text for p in docx.Document(io.BytesIO(word.content)).paragraphs)
        for expected in ("Battle test: Dana Whitfield v. Summit Peak Roofing LLC", "not legal advice", "Complaint",
                         "Opposition to Motion for Summary Judgment", "Utah Code § 99-99-999",
                         "[⚠ NOT FOUND: no such section in the law index]", "[CITATION NEEDED:",
                         "Authorities quoted to the models", "https://le.utah.gov/"):
            with self.subTest(expected=expected):
                self.assertIn(expected, text)

        pdf = self.client.get(f"{case_url}/download.pdf")
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf.headers["content-type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF-"))
        # Cover and citation check, one page per document at least, and the appendix.
        self.assertGreaterEqual(len(re.findall(rb"/Type\s*/Page\b", pdf.content)), 5)

        self.assertEqual(self.client.get(f"{case_url}/download.exe").status_code, 404)
        # Nothing extra is stored with the case.
        folder = self.app.state.store.root / case_url.rsplit("/", 1)[1]
        self.assertEqual(sorted(f.name for f in folder.iterdir()), ["input.md", "result.json", "result.md"])

    def test_no_word_or_pdf_before_a_case_has_finished(self):
        case_url = self.finished_case()
        self.app.state.store.update(case_url.rsplit("/", 1)[1], status="failed", error="boom")
        for fmt in ("md", "docx", "pdf"):
            self.assertEqual(self.client.get(f"{case_url}/download.{fmt}").status_code, 404)

    def test_case_page_lists_own_cases_newest_first(self):
        first = self.finished_case(plaintiff="Avery First")
        second = self.finished_case(plaintiff="Blake Second")
        failed = self.finished_case(plaintiff="Casey Third")
        self.app.state.store.update(failed.rsplit("/", 1)[1], status="failed", error="boom")
        other = TestClient(self.app)
        self.login("other", client=other)
        theirs = self.submit(client=other, plaintiff="Someone Else").headers["location"]
        self.assertTrue(self.app.state.worker.wait_idle())

        for url in (first, failed):  # a results page and a failure page
            with self.subTest(url=url):
                nav = re.search(r'<nav class="case-nav".*?</nav>', self.client.get(url).text, re.DOTALL).group()
                self.assertEqual(re.findall(r'href="(/cases/\w+)"', nav), [failed, second, first])
                self.assertEqual(re.findall(r'href="(/cases/\w+)" aria-current="page"', nav), [url])
                self.assertNotIn("Someone Else", nav)
        self.assertNotIn(theirs, self.client.get(first).text)
        # The new-case form lists them too, with none marked as open.
        for page in (self.client.get("/cases/new").text, self.submit(state="XX").text):
            self.assertEqual(re.findall(r'href="(/cases/[0-9a-f]{32})"', page), [failed, second, first])
            self.assertNotIn("aria-current", page)
        self.assertNotIn("case-nav", other.get("/account").text)

    def test_delete_case(self):
        case_url = self.finished_case()
        job_id = case_url.rsplit("/", 1)[1]
        page = self.client.get(case_url).text
        self.assertIn(f'href="{case_url}/delete"', page)
        self.assertRegex(page, r"deleted automatically on \d{4}-\d{2}-\d{2}")
        self.assertIn("Delete this case?", self.client.get(f"{case_url}/delete").text)

        self.assertEqual(self.client.post(f"{case_url}/delete", data={"csrf": "wrong"}).status_code, 403)
        other = TestClient(self.app)
        self.login("other", client=other)
        self.assertEqual(other.post(f"{case_url}/delete", data={"csrf": self.csrf(other)}).status_code, 404)
        self.assertIsNotNone(self.app.state.store.get(job_id))  # neither attempt deleted it

        response = self.client.post(f"{case_url}/delete", data={"csrf": self.csrf()}, follow_redirects=False)
        self.assertEqual(response.headers["location"], "/?deleted=1")
        self.assertIsNone(self.app.state.store.get(job_id))
        self.assertEqual(self.client.get(case_url).status_code, 404)
        self.assertIn("Case deleted.", self.client.get("/?deleted=1").text)

    def test_running_case_cannot_be_deleted(self):
        job = self.app.state.store.create(self.app.state.auth.get_user("dana").id, "UT", "facts", "x", 1, "t")
        self.app.state.store.update(job.id, status="running")
        response = self.client.post(f"/cases/{job.id}/delete", data={"csrf": self.csrf()})
        self.assertEqual(response.status_code, 409)
        self.assertIsNotNone(self.app.state.store.get(job.id))
        self.app.state.store.update(job.id, status="failed")  # let the worker skip it

    def test_case_list_shows_retention(self):
        self.finished_case()
        page = self.client.get("/").text
        self.assertIn("deleted automatically 90 days after", page)
        self.assertRegex(page, r'<td class="muted">\d{4}-\d{2}-\d{2}</td>')

    def test_events_stream_stages_and_finish(self):
        case_url = self.finished_case()
        body = self.client.get(f"{case_url}/events").text
        stages = re.findall(r'"title": "([^"]+)"', body)
        self.assertEqual(stages[:3], ["Legal research", "Selecting authorities", "Complaint"])
        self.assertIn("event: text", body)
        self.assertTrue(body.rstrip().endswith('"error": ""}'))  # the final "done" event

    def test_validation_errors_keep_the_form(self):
        response = self.submit(state="NV", plaintiff="Kept Name", timeline="")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Choose a state.", response.text)
        self.assertIn("Fill in at least the plaintiff and what happened.", response.text)
        self.assertIn("Kept Name", response.text)
        self.assertIn("Start your first case", self.client.get("/").text)

    def test_unknown_case_is_404(self):
        self.assertEqual(self.client.get("/cases/nope").status_code, 404)
        self.assertEqual(self.client.get("/cases/nope/download.md").status_code, 404)


@unittest.skipIf(TestClient is None, "web dependencies not installed")
class RenderTest(unittest.TestCase):
    def doc(self, text, checks=()):
        return {"text": text, "checks": list(checks), "unchecked": [], "authorities": []}

    def test_model_output_is_escaped(self):
        html = document_html(self.doc('<script>alert(1)</script> **bold**'))
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)
        self.assertIn("<strong>bold</strong>", html)

    def test_links_longest_citation_first(self):
        section = lambda c: {"citation": c, "source_url": f"https://x.gov/{c[-2:]}"}  # noqa: E731
        html = document_html(self.doc(
            "See § 1-1-12 and § 1-1-1.",
            [{"text": "§ 1-1-1", "status": "verified", "section": section("§ 1-1-1")},
             {"text": "§ 1-1-12", "status": "verified", "section": section("§ 1-1-12")}],
        ))
        self.assertIn('href="https://x.gov/12"', html)
        self.assertIn('href="https://x.gov/-1"', html)
        self.assertEqual(html.count("<a "), 2)

    def test_doubted_citations_are_listed_with_the_checkers_note(self):
        section = {"citation": "Utah Code § 1-1-1", "source_url": "https://x.gov/1"}
        doc = self.doc("Venue lies under § 1-1-1 [⚠ MAY NOT SUPPORT THIS].", [
            {"text": "§ 1-1-1", "status": "verified", "section": section, "support": "does_not_support",
             "support_reason": "It is about <fees>."},
            {"text": "§ 1-1-1", "status": "in_force_not_provided", "section": section, "support": "supports",
             "support_reason": "ok"}])
        self.assertEqual(render.check_summary(doc, checked=True)["doubted"], 1)
        self.assertEqual(render.check_summary(doc)["doubted"], "–")
        doubt, unseen = render.notable_checks(doc)
        self.assertEqual((doubt["label"], doubt["reason"], doubt["css"]),
                         ("⚠ may not support what it's cited for", "It is about <fees>.", "warn"))
        self.assertEqual(unseen["reason"], "")
        self.assertIn('<mark class="flag">[⚠ MAY NOT SUPPORT THIS]</mark>', document_html(doc))
        # Results saved before the check existed have no verdicts and no checker model.
        old = self.doc("x", [{"text": "§ 1-1-1", "status": "verified", "section": section}])
        self.assertEqual(render.notable_checks(old), [])
        self.assertFalse(render.support_checked({"documents": [old]}))

        result = {**ExportTest.RESULT, "checker_model": "c", "documents": [
            {**ExportTest.RESULT["documents"][0], **doc}]}
        blocks = export.blocks(result, "A v. B", "now")
        texts = [t for b in blocks for t in [b.text, *(i[0] for i in b.items if b.kind == "bullets")]]
        self.assertIn("Citation checker model: c", texts)
        self.assertIn("⚠ may not support what it's cited for: § 1-1-1. Checker's note: It is about <fees>.", texts)
        self.assertTrue(any("A model also compared each one" in t for t in texts))
        self.assertTrue(export.to_pdf(result, "A v. B", "now").startswith(b"%PDF-"))

    def test_case_title_uses_names_only(self):
        self.assertEqual(case_title({"plaintiff": "Dana Whitfield, homeowner in Provo",
                                     "defendant": "Summit Peak Roofing LLC (Orem)"}),
                         "Dana Whitfield v. Summit Peak Roofing LLC")
        self.assertEqual(case_title({"plaintiff": "Dana Whitfield"}), "Dana Whitfield")

    def test_facts_text_skips_empty_fields(self):
        text = facts_text({"plaintiff": "A", "defendant": " ", "timeline": "B"})
        self.assertIn("- **Your client (the plaintiff):** A", text)
        self.assertNotIn("defendant", text)


if __name__ == "__main__":
    unittest.main()


@unittest.skipIf(TestClient is None, "web dependencies not installed")
class ExportTest(unittest.TestCase):
    RESULT = {
        "state": "Utah", "rounds": 1, "plaintiff_model": "m", "defendant_model": "m", "law_meta": {},
        "research": {"plaintiff": ["venue"]},
        "documents": [{
            "title": "Complaint", "role": "plaintiff", "generated": True, "unchecked": [], "authorities": [],
            "checks": [{"text": "Utah Code § 1-1-1", "status": "not_found", "section": None}],
            "text": "**CAPTION**\nLine two <b>&amp;</b>\n\nSee Utah Code § 1-1-1 [⚠ NOT FOUND: x] and "
                    "[FACT NEEDED: the date]. Nguyễn → 日本",
        }],
    }

    def test_spans_split_marks_placeholders_and_bold(self):
        self.assertEqual(
            export.spans("a **b** [⚠ NOT FOUND: x] c [CITATION NEEDED: y]"),
            [("a ", ""), ("b", "bold"), (" ", ""), ("[⚠ NOT FOUND: x]", "flag"), (" c ", ""),
             ("[CITATION NEEDED: y]", "placeholder")])

    def test_pdf_text_keeps_legal_symbols_and_degrades_the_rest(self):
        self.assertEqual(export.pdf_text("§ 78B-2-309 — “quoted”…"), "§ 78B-2-309 — “quoted”…")
        self.assertEqual(export.pdf_text("[⚠ NOT FOUND] Nguyễn → 日"), "[! NOT FOUND] Nguyen -> ?")

    def test_drafts_keep_their_paragraphs_and_are_never_markup(self):
        blocks = export.blocks(self.RESULT, "A v. B", "2026-10-05 20:00")
        paras = [b.text for b in blocks if b.kind == "para"]
        self.assertIn("**CAPTION**\nLine two <b>&amp;</b>", paras)
        # Both formats build from awkward text without error, and the PDF
        # treats tags in a draft as text.
        self.assertIn("&lt;b&gt;&amp;amp;&lt;/b&gt;", export._markup("Line two <b>&amp;</b>"))
        self.assertTrue(export.to_pdf(self.RESULT, "A v. B <i>", "now").startswith(b"%PDF-"))
        self.assertTrue(export.to_docx(self.RESULT, "A v. B <i>", "now").startswith(b"PK"))

    def test_possible_references_are_the_uncited_authorities(self):
        cited = {"citation": "Utah R. Civ. P. 56", "section_title": "Rule 56", "text": "A party may move.",
                 "source_url": "https://x.gov/56"}
        spare = {"citation": "Utah Code § 78B-2-309", "section_title": "Within six years",
                 "text": "An action  may be\nbrought " * 40, "source_url": ""}
        doc = {**self.RESULT["documents"][0], "authorities": [cited, spare],
               "checks": [{"text": "Rule 56", "status": "verified", "section": cited}]}
        result = {**self.RESULT, "documents": [doc, {**doc, "title": "Motion", "checks": []}]}

        (possible,) = render.possible_references(result)
        self.assertEqual(possible["citation"], "Utah Code § 78B-2-309")
        self.assertTrue(possible["excerpt"].startswith("An action may be brought An action"))
        self.assertTrue(possible["excerpt"].endswith(" …"))
        self.assertLess(len(possible["excerpt"]), 310)

        blocks = export.blocks(result, "A v. B", "now")
        at = [b.text for b in blocks].index("Possible references")
        (bullet,) = blocks[at + 2].items
        self.assertTrue(bullet[0].startswith("**Utah Code § 78B-2-309** Within six years — An action"))
        self.assertTrue(bullet[0].endswith("⚠ no official source URL"))
        self.assertTrue(export.to_pdf(result, "A v. B", "now").startswith(b"%PDF-"))
        self.assertTrue(export.to_docx(result, "A v. B", "now").startswith(b"PK"))
        # The fit check's verdicts: unsure joins the possible references, does-not-apply the appendix.
        maybe = {**spare, "citation": "Utah Code § 15-1-1", "text": "Interest."}
        wrong = {**spare, "citation": "Utah Code § 75-1-303", "section_title": "Venue", "text": "Probate."}
        fitted = {**result, "fit": {"plaintiff": [
            {"section": cited, "verdict": "applies", "reason": "fits"},
            {"section": maybe, "verdict": "unsure", "reason": "Depends on the contract."},
            {"section": wrong, "verdict": "does_not_apply", "reason": ""}]}}
        held = render.possible_references(fitted)[1]
        self.assertEqual((held["citation"], held["held_back"], held["excerpt"]),
                         ("Utah Code § 15-1-1", "Depends on the contract.", "Interest."))
        self.assertEqual([(a["citation"], a["reason"]) for a in render.set_aside(fitted)],
                         [("Utah Code § 75-1-303", "none given")])
        texts = [i[0] for b in export.blocks(fitted, "A v. B", "now") if b.kind == "bullets" for i in b.items]
        self.assertTrue(any("§ 15-1-1" in t and "Checker's note: Depends on the contract." in t for t in texts))
        self.assertIn("Utah Code § 75-1-303 Venue. Checker's note: none given", texts)
        self.assertEqual(render.set_aside(result), [])  # saved before the fit check existed
        # Nothing given and uncited: no section at all.
        self.assertNotIn("Possible references", [b.text for b in export.blocks(self.RESULT, "A v. B", "now")])
